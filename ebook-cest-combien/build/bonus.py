"""Génère les bonus à partir du manuscrit : scripts et prompts en texte brut,
tableau de suivi Excel et kit imprimable PDF.

Usage : python3 build/bonus.py   (depuis le dossier ebook-cest-combien)
"""
import html, re
from pathlib import Path

import markdown
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent.parent
MS, BUILD, OUT = ROOT / "manuscrit", ROOT / "build", ROOT / "livrables" / "bonus"
CHAPTERS = [MS / f"ch{i:02d}.md" for i in range(1, 13)]
HEADER = """« C'EST COMBIEN ? » — {title}
Bonus du livre d'Amadou Traoré · DÉO-CI · © 2026
Usage personnel de l'acheteur. Copiez, adaptez, remplacez ce qui est entre [crochets].
{rule}
"""


def plain(s):
    s = re.sub(r"<br\s*/?>", "\n", s)
    s = re.sub(r'<span class="k">(.*?)</span>', lambda m: "\n" + m.group(1).upper() + "\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = s.replace("**", "").replace("*", "")
    s = html.unescape(s)
    lines = [ln.rstrip() for ln in s.strip().splitlines()]
    out, blank = [], False
    for ln in lines:
        if not ln.strip():
            if not blank:
                out.append("")
            blank = True
        else:
            out.append(ln.strip())
            blank = False
    return "\n".join(out).strip()


def chapter_title(src):
    num = re.search(r'<p class="chap-num">(Chapitre \d+)', src).group(1)
    title = re.search(r'<h1 class="chap">(.*?)</h1>', src).group(1)
    return f"{num} — {html.unescape(title)}"


def blocks(src, cls):
    """Renvoie (titre, contenu) des blocs <div class="cls" ...> du fichier."""
    res = []
    for m in re.finditer(rf'<div class="{cls}"[^>]*>', src):
        start, depth, i = m.end(), 1, m.end()
        while depth:
            o = src.find("<div", i)
            c = src.find("</div>", i)
            if o != -1 and o < c:
                depth += 1
                i = o + 4
            else:
                depth -= 1
                i = c + 6
        body = src[start:i - 6]
        t = re.search(r'<(?:span|div) class="(?:sh|ph)">(.*?)</(?:span|div)>', body)
        title = plain(t.group(1)) if t else ""
        if t:
            body = body.replace(t.group(0), "", 1)
        res.append((title, plain(body)))
    return res


def write_txt(name, title, cls):
    rule = "=" * 64
    parts = [HEADER.format(title=title, rule=rule)]
    n = 0
    for ch in CHAPTERS:
        src = ch.read_text(encoding="utf-8")
        found = blocks(src, cls)
        if not found:
            continue
        parts.append(f"\n{rule}\n{chapter_title(src).upper()}\n{rule}\n")
        for t, body in found:
            n += 1
            parts.append(f"\n--- {t or 'Modèle'} ---\n{body}\n")
    (OUT / name).write_text("\n".join(parts), encoding="utf-8")
    return n


def write_xlsx():
    wb = Workbook()
    green, orange = "0E5A47", "E07A2E"
    thin = Side(style="thin", color="D9E2DE")
    sheets = {
        "Commandes": (["Date", "Client (nom enregistré)", "Produit", "Quantité", "Prix (FCFA)", "Frais livraison (FCFA)", "Total (FCFA)",
                       "Zone / ville", "Mode de paiement", "Statut paiement", "Statut livraison", "Livreur", "Nouveau ou ancien client", "Remarque"],
                      [12, 24, 26, 9, 12, 14, 12, 16, 16, 15, 15, 14, 16, 30]),
        "Demandes et relances": (["Date demande", "Client", "Produit demandé", "Étiquette", "Réponse en 3 temps envoyée ?", "Objection", "J+1 envoyé", "J+3 envoyé",
                                  "J+7 envoyé", "Résultat", "Remarque"], [13, 22, 24, 18, 16, 22, 11, 11, 11, 16, 30]),
        "Bilan hebdo": (["Semaine du", "Vues moyennes statuts", "Demandes reçues", "Commandes payées", "Commandes / demandes", "Clients revenus",
                         "Ce qui a marché", "Étape qui fuit le plus", "Amélioration de la semaine suivante"], [13, 14, 13, 13, 14, 12, 30, 18, 34]),
        "Clients fidèles": (["Client", "Téléphone", "Commune", "Nb d'achats", "Dernier achat", "Préférences (dites par le client)", "Avantage fidélité dû",
                             "Accord diffusion (OUI/STOP)", "Parrainé par"], [22, 16, 14, 11, 13, 32, 18, 16, 18]),
        "Objections": (["Objection", "Nb de fois", "Réponse rapide", "Preuve à envoyer", "Option alternative", "Ce que ça révèle sur mon offre"],
                       [26, 10, 16, 26, 24, 34]),
    }
    first = True
    for name, (cols, widths) in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws.append(cols)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
            c = ws.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=green if name != "Bilan hebdo" else orange)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            c.border = Border(bottom=thin)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
    ws = wb["Commandes"]
    for r in range(2, 501):
        ws[f"G{r}"] = f'=IF(OR(E{r}="",D{r}=""),"",D{r}*E{r}+N(F{r}))'
        for col in "EFG":
            ws[f"{col}{r}"].number_format = '# ##0'
    for col, opts in {"I": "Espèces,Wave,Orange Money,MTN MoMo,Moov Money,Autre", "J": "Payé,Acompte reçu,À payer à la livraison,Non payé",
                      "K": "À livrer,En route,Livrée,Injoignable,Refusée,Expédiée", "M": "Nouveau,Ancien"}.items():
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}500")
    ws = wb["Demandes et relances"]
    dv = DataValidation(type="list", formula1='"Nouveau contact,A demandé le prix,Commande en cours,Payé,Client fidèle"', allow_blank=True)
    ws.add_data_validation(dv); dv.add("D2:D500")
    dv2 = DataValidation(type="list", formula1='"Commande,Pas de réponse,Non,En attente"', allow_blank=True)
    ws.add_data_validation(dv2); dv2.add("J2:J500")
    ws = wb["Bilan hebdo"]
    for r in range(2, 54):
        ws[f"E{r}"] = f'=IF(OR(C{r}="",C{r}=0),"",D{r}/C{r})'
        ws[f"E{r}"].number_format = "0%"
    ws = wb["Objections"]
    for row in [("C'est cher", "", "/cher", "Photo de près, détail de la matière", "Modèle moins cher, format découverte", ""),
                ("Je vais réfléchir", "", "/reflechir", "", "Réservation 24 h", ""),
                ("Payer à la livraison ?", "", "/paiement", "Avis clients, vidéo des colis", "Acompte réduit", ""),
                ("C'est original ?", "", "/original", "Vidéo, emballage, facture", "", ""),
                ("Vous faites un prix ?", "", "/prix-fixe", "", "Avantage structuré", "")]:
        ws.append(row)
    wb.save(OUT / "Tableau-de-suivi_C-est-combien.xlsx")


def section(path, start, end=None):
    s = (MS / path).read_text(encoding="utf-8")
    i = s.index(start)
    j = s.index(end, i + len(start)) if end else len(s)
    return s[i:j]


def write_kit():
    parts = [
        '<div class="title-page"><p class="tp-brand">DÉO-CI · BONUS</p><p class="tp-title">Le kit imprimable</p>'
        '<p class="tp-sub">Les fiches, checklists et grilles de « C\'est combien ? », à imprimer ou à garder sur votre téléphone.</p>'
        '<p class="tp-author">Amadou Traoré</p></div>',
        '<div class="front-chapter" markdown="1">\n<h1 class="front">1. L\'autodiagnostic</h1>\n\n' + section("ch01.md", "Répondez honnêtement", "<div class=\"box retenir\" markdown=\"1\">\n**Rendez-vous") + "\n</div>",
        '<div class="front-chapter" markdown="1">\n<h1 class="front">2. Configuration WhatsApp Business</h1>\n\n' + section("ch02.md", '<div class="cards c2" markdown="1">', "### 2. Trois modèles") + "\n</div>",
        '<div class="front-chapter" markdown="1">\n<h1 class="front">3. Statuts : semaine type et 30 jours</h1>\n\n'
        + section("ch05.md", "| Jour | Thème principal | En plus |", "Un restaurant publiera")
        + section("ch05.md", "### 1. Calendrier de statuts sur 30 jours", "<p class=\"src\">") + "\n</div>",
        '<div class="front-chapter" markdown="1">\n<h1 class="front">4. Les deux méthodes clés</h1>\n\n### La réponse au prix en 3 temps\n\n{{inc:trois-temps}}\n\n### La méthode ACRA face aux objections\n\n{{inc:acra}}\n\n'
        '**A**ccueillir · **C**larifier · **R**épondre · **A**vancer\n\n### Le calendrier de relance\n\n{{inc:relance}}\n\n### Les étiquettes\n\n{{inc:etiquettes}}\n</div>',
        '<div class="front-chapter" markdown="1">\n<h1 class="front">5. Paiement et anti-arnaque</h1>\n\n' + section("ch09.md", "### 2. Fiche du livreur", None) + "\n</div>",
        '<div class="front-chapter" markdown="1">\n<h1 class="front">6. Routine et bilan</h1>\n\n' + section("ch12.md", "### 1. Checklist de la routine quotidienne", None)
        + "\n\n" + section("ch12.md", "**Exercice 2 : votre premier bilan hebdomadaire**", "**Exercice 3") + "\n</div>",
        section("99-fin.md", '<div class="front-chapter plan" markdown="1">', '<div class="front-chapter" markdown="1">\n<h1 class="front">La boîte'),
    ]
    src = "\n\n".join(parts).replace('<p class="mk">QQplan30QQ</p>', "")
    src = re.sub(r"\{\{inc:([\w-]+)\}\}", lambda m: (BUILD / "snippets" / f"{m.group(1)}.html").read_text(encoding="utf-8"), src)
    body = markdown.markdown(src, extensions=["tables", "md_in_html", "attr_list", "sane_lists", "nl2br"])
    css = (BUILD / "style.css").read_text(encoding="utf-8").replace("« C'est combien ? »  ·  DÉO-CI", "Kit imprimable · « C'est combien ? » · DÉO-CI")
    page = BUILD / "kit.html"
    page.write_text(f'<!doctype html><html lang="fr"><head><meta charset="utf-8"><title>Kit imprimable</title><style>{css}</style></head><body>{body}</body></html>', encoding="utf-8")
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium")
        p = b.new_page()
        p.goto(page.as_uri()); p.wait_for_load_state("networkidle"); p.evaluate("document.fonts.ready")
        p.pdf(path=str(OUT / "Kit-imprimable_C-est-combien.pdf"), prefer_css_page_size=True, print_background=True)
        b.close()


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    print("scripts :", write_txt("Scripts-a-copier_C-est-combien.txt", "LES SCRIPTS À COPIER", "script"))
    print("prompts :", write_txt("Prompts-IA_C-est-combien.txt", "LES PROMPTS IA", "prompt"))
    write_xlsx(); print("xlsx ok")
    write_kit(); print("kit ok")
